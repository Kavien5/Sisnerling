import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palet Warna Profesional & Modern
    C_BG_DARK = RGBColor(15, 23, 42)      # #0f172a (Navy Slate)
    C_CARD_BG = RGBColor(30, 41, 59)      # #1e293b (Card)
    C_CARD_ALT = RGBColor(24, 33, 47)     # #18212f
    C_EMERALD = RGBColor(16, 185, 129)    # #10b981 (Accent Green)
    C_BLUE = RGBColor(59, 130, 246)       # #3b82f6 (Accent Blue)
    C_AMBER = RGBColor(245, 158, 11)      # #f59e0b (Accent Warning/Gold)
    C_WHITE = RGBColor(255, 255, 255)
    C_GRAY_LIGHT = RGBColor(226, 232, 240)# #e2e8f0 (Text Body)
    C_GRAY_MUTED = RGBColor(148, 163, 184)# #94a3b8

    def add_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SISNERLING — PRESENTASI SISTEM APLIKASI WEB"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_EMERALD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.8), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(25)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def set_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ==========================================
    # SLIDE 1: JUDUL
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s1)

    # Accent bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.8), Inches(0.12))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = C_EMERALD
    bar1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "SISTEM PENGOLAHAN DATA BERBASIS WEB"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_EMERALD
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "SISNERLING"
    p_main.font.size = Pt(54)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "(Sistem Neraca Lingkungan)"
    p_sub.font.size = Pt(24)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_BLUE
    p_sub.space_after = Pt(14)

    p_tag = tf1.add_paragraph()
    p_tag.text = "Digitalisasi Pengelolaan Data Lingkungan Tahunan untuk Pengambilan Keputusan Strategis"
    p_tag.font.size = Pt(18)
    p_tag.font.color.rgb = C_GRAY_LIGHT
    p_tag.space_after = Pt(16)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Platform terpusat untuk menginput, memvalidasi, merekap, dan menganalisis neraca sumber daya alam secara berkala."
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = C_GRAY_MUTED

    set_speaker_notes(s1, 
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Selamat pagi/siang bapak/ibu sekalian. Pada kesempatan ini saya ingin memperkenalkan SISNERLING (Sistem Neraca Lingkungan), "
        "sebuah platform aplikasi berbasis web yang kami kembangkan untuk mendigitalisasi seluruh siklus pengelolaan data lingkungan tahunan. "
        "Sistem ini hadir untuk menjembatani kebutuhan data lingkungan yang akurat, cepat, dan terpercaya sebagai fondasi pengambilan keputusan dan kebijakan strategis."
    )

    # ==========================================
    # SLIDE 2: LATAR BELAKANG & MASALAH
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s2)
    add_header(s2, "Latar Belakang & Kendala Pengelolaan Data Manual")

    problems = [
        ("01", "Data Tercecer & Desentralisasi", 
         "Data lingkungan tersimpan di puluhan file Excel terpisah di komputer masing-masing staf. Sulit melacak versi data terkini dan rentan kehilangan data saat terjadi pergantian personil."),
        ("02", "Rekapitulasi Lambat & Rumit", 
         "Penyusunan neraca membutuhkan proses kompilasi manual, penyalinan berulang, dan perhitungan lintas-file yang memakan waktu berminggu-minggu hingga berbulan-bulan."),
        ("03", "Ketidaksesuaian & Ketidaksinkronan Data", 
         "Rumus formula Excel yang rawan rusak (#REF!, #VALUE!), perbedaan asumsi satuan, serta tidak adanya validasi otomatis menghasilkan inkonsistensi antar instansi/seksi.")
    ]

    for i, (num, p_title, p_desc) in enumerate(problems):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        w = Inches(3.6)
        h = Inches(4.8)
        add_card(s2, left, top, w, h, C_CARD_BG, border_color=C_AMBER if i==2 else None)

        tb = s2.shapes.add_textbox(left + Inches(0.25), top + Inches(0.3), w - Inches(0.5), h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"KENDALA {num}"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = C_AMBER
        p_num.space_after = Pt(12)

        p_t = tf.add_paragraph()
        p_t.text = p_title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = p_desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = C_GRAY_LIGHT

    set_speaker_notes(s2,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Bapak/Ibu, mari kita lihat kondisi sebelum adanya SISNERLING. Sebelumnya, pengelolaan data neraca lingkungan masih bertumpu pada spreadsheet Excel manual. "
        "Hal ini menimbulkan tiga masalah besar:\n"
        "1. Data tercecer di berbagai unit dan komputer individu tanpa satu pusat repositori.\n"
        "2. Rekapitulasi sangat lambat karena harus mencocokkan puluhan file berbeda secara manual setiap tahun.\n"
        "3. Ketidaksesuaian data akibat rumus yang rusak atau kesalahan pengetikan manusia yang baru disadari saat laporan dicetak. Kendala-kendala inilah yang kami selesaikan."
    )

    # ==========================================
    # SLIDE 3: SOLUSI & PENGENALAN SISTEM
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s3)
    add_header(s3, "Solusi & Pengenalan SISNERLING")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), C_CARD_BG, border_color=C_EMERALD)
    tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p_lt = tf_l.paragraphs[0]
    p_lt.text = "Apa itu SISNERLING?"
    p_lt.font.size = Pt(20)
    p_lt.font.bold = True
    p_lt.font.color.rgb = C_EMERALD
    p_lt.space_after = Pt(12)

    l_items = [
        "Aplikasi web terintegrasi untuk mencatat, mengolah, dan mengevaluasi data neraca lingkungan tahunan secara berkala.",
        "Mengadopsi antarmuka spreadsheet interaktif yang mudah dipelajari, didukung database terpusat yang aman.",
        "Menggantikan tumpukan file fisik dan Excel yang rapuh menjadi satu pintu akses resmi terpadu.",
        "Menyediakan sistem formulasi otomatis sehingga angka neraca fisik dan moneter langsung terhitung konsisten."
    ]
    for it in l_items:
        p = tf_l.add_paragraph()
        p.text = f"✔  {it}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_GRAY_LIGHT
        p.space_after = Pt(10)

    add_card(s3, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), C_CARD_BG, border_color=C_BLUE)
    tb_r = s3.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p_rt = tf_r.paragraphs[0]
    p_rt.text = "Tujuan Utama Sistem"
    p_rt.font.size = Pt(20)
    p_rt.font.bold = True
    p_rt.font.color.rgb = C_BLUE
    p_rt.space_after = Pt(12)

    r_items = [
        "Digitalisasi Penuh: Mengeliminasi alur kerja manual berbasis kertas dan spreadsheet yang tersebar.",
        "Sentralisasi Data: Mewujudkan satu sumber kebenaran data lingkungan hidup yang terverifikasi.",
        "Dukungan Kebijakan Berbasis Bukti: Menyajikan laporan analitis cepat untuk pimpinan dalam merumuskan kebijakan pelestarian alam.",
        "Transparansi & Akuntabilitas: Menyimpan riwayat perubahan nilai dan aktor pengubah secara jelas (Audit Trail)."
    ]
    for it in r_items:
        p = tf_r.add_paragraph()
        p.text = f"★  {it}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_GRAY_LIGHT
        p.space_after = Pt(10)

    set_speaker_notes(s3,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "SISNERLING hadir sebagai jawaban atas tantangan tersebut. Sistem ini mengintegrasikan seluruh tahapan pengelolaan neraca lingkungan ke dalam sebuah portal web terpusat. "
        "Tujuan utamanya jelas: kita ingin seluruh proses digital, data tidak lagi tercecer, dan pimpinan mendapatkan angka neraca lingkungan yang akurat kapan pun dibutuhkan untuk mengambil keputusan strategis berbasis data riil."
    )

    # ==========================================
    # SLIDE 4: FITUR UTAMA WEB
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s4)
    add_header(s4, "Fitur-Fitur Utama Aplikasi Web SISNERLING")

    features = [
        ("1. Manajemen Pengguna Multi-Role", 
         "Pemisahan hak akses ketat antara Operator (entri data penugasan), Validator (pemeriksa dan penyetujui data), dan Admin/Super Admin (pengelola sistem & konfigurasi penuh)."),
        ("2. Form Input Tahunan Interaktif", 
         "Tampilan formulir berbasis grid spreadsheet yang familiar, dilengkapi validasi tipe angka otomatis, deteksi kesalahan langsung, dan penyimpanan berkala yang aman."),
        ("3. Mekanisme Validasi Berjenjang", 
         "Alur verifikasi data sebelum disahkan: status draf/pending, verifikasi kelayakan oleh validator, hingga penguncian data (freeze/lock) saat periode ditutup."),
        ("4. Dashboard Analytics & Export Laporan", 
         "Visualisasi tren data deret waktu tahunan, pemantauan saldo stok (awal, pertambahan, deplesi, akhir), serta export laporan instan ke format Excel dan PDF siap cetak.")
    ]

    for i, (ft, fd) in enumerate(features):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.5)
        w = Inches(5.6)
        h = Inches(2.25)

        add_card(s4, left, top, w, h, C_CARD_BG)
        tb = s4.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = ft
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = C_EMERALD
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = fd
        p2.font.size = Pt(12.5)
        p2.font.color.rgb = C_GRAY_LIGHT

    set_speaker_notes(s4,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Ada empat pilar fitur utama di SISNERLING:\n"
        "Pertama, Manajemen User Multi-Role: membagi peran dengan jelas antara Operator pengisi data, Validator pemeriksa, dan Admin.\n"
        "Kedua, Form Input Interaktif: nyaman digunakan seperti spreadsheet tetapi dilengkapi validasi sistem agar tidak ada salah format.\n"
        "Ketiga, Validasi Berjenjang: data tidak bisa langsung masuk ke laporan sebelum diverifikasi oleh validator yang berwenang.\n"
        "Keempat, Dashboard Analytics & Export: menyajikan grafik ringkas dan kemampuan mengunduh laporan neraca dalam sekejap."
    )

    # ==========================================
    # SLIDE 5: ALUR KERJA (WORKFLOW)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s5)
    add_header(s5, "Alur Kerja Sistem (Workflow Operasional)")

    steps = [
        ("TAHAP 01", "Login & Hak Akses", "Pengguna masuk sesuai kredensial dan role (Operator / Validator / Admin) dengan enkripsi aman."),
        ("TAHAP 02", "Pengisian Data", "Operator mengisi data tahunan pada form spreadsheet sesuai lingkup tabel yang ditugaskan."),
        ("TAHAP 03", "Kalkulasi Otomatis", "Sistem menjalankan perhitungan formula neraca secara instan tanpa perlu ketik rumus manual."),
        ("TAHAP 04", "Proses Validasi", "Validator memeriksa kelengkapan & kewajaran data. Dapat menyetujui atau meminta revisi."),
        ("TAHAP 05", "Publikasi & Analisis", "Data resmi terkunci, langsung tampil di Dashboard Analytics dan siap diekspor menjadi laporan.")
    ]

    step_w = Inches(2.25)
    step_h = Inches(4.8)
    step_top = Inches(1.8)

    for i, (stg, st_title, st_desc) in enumerate(steps):
        left = Inches(0.8 + i * 2.45)
        border = C_EMERALD if i == 3 else (C_BLUE if i == 4 else None)
        add_card(s5, left, step_top, step_w, step_h, C_CARD_BG, border_color=border)

        tb = s5.shapes.add_textbox(left + Inches(0.15), step_top + Inches(0.25), step_w - Inches(0.3), step_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = stg
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = C_AMBER
        p0.space_after = Pt(10)

        p1 = tf.add_paragraph()
        p1.text = st_title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_GRAY_LIGHT

    set_speaker_notes(s5,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Bagaimana cara sistem ini bekerja dalam operasional harian? Alurnya sangat runtut dan tertata:\n"
        "Dimulai dari Operator yang login dan memasukkan data tahunan pada tabel yang menjadi tanggung jawabnya. "
        "Secara otomatis sistem melakukan perhitungan formula neraca. Kemudian data tersebut diajukan ke Validator untuk diperiksa. "
        "Setelah disetujui, data akan dipublikasikan ke Dashboard untuk dianalisis oleh pimpinan dan diekspor sebagai dokumen resmi. Tidak ada lagi lompatan proses yang tidak tercatat."
    )

    # ==========================================
    # SLIDE 6: MENGAPA MEMBUTUHKAN SISNERLING
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s6)
    add_header(s6, "Mengapa Kita Membutuhkan SISNERLING?")

    benefits = [
        ("Akurasi Data Terjamin", 
         "Meniadakan kesalahan manusia (*human error*) dan formula Excel yang terhapus secara tidak sengaja melalui standarisasi validasi input dan rumus otomatis."),
        ("Efisiensi Waktu Signifikan", 
         "Memotong waktu kompilasi dan rekapitulasi data dari berminggu-minggu menjadi instan (hitungan detik) begitu data dasar tervalidasi."),
        ("Keamanan & Histori Lengkap", 
         "Merekam setiap riwayat input, revisi, dan aktor pengubah (*Audit Trail*), sehingga data memiliki transparansi tinggi dan siap diaudit."),
        ("Single Source of Truth", 
         "Menyediakan satu-satunya sumber data resmi yang konsisten bagi seluruh pemangku kepentingan, menghilangkan kebingungan perbedaan versi dokumen.")
    ]

    for i, (bt, bd) in enumerate(benefits):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.5)
        w = Inches(5.6)
        h = Inches(2.25)

        add_card(s6, left, top, w, h, C_CARD_BG, border_color=C_BLUE if (i%2==1) else C_EMERALD)
        tb = s6.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f"★  {bt}"
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = bd
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_GRAY_LIGHT

    set_speaker_notes(s6,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Mengapa investasi pada sistem SISNERLING ini krusial bagi organisasi kita? "
        "Ada empat nilai strategis yang kita dapatkan:\n"
        "1. Akurasi: angka yang disajikan terjamin karena rumus dijalankan terpusat oleh sistem.\n"
        "2. Efisiensi waktu: kita menghemat waktu kerja pegawai dari merekap manual menjadi fokus pada analisis data.\n"
        "3. Keamanan data: ada histori lengkap siapa yang mengubah apa dan kapan.\n"
        "4. Single Source of Truth: tidak ada lagi perdebatan versi laporan mana yang paling update karena semua mengacu pada SISNERLING."
    )

    # ==========================================
    # SLIDE 7: KESIMPULAN & ROADMAP
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s7)
    add_header(s7, "Kesimpulan & Rencana Pengembangan (Roadmap)")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), C_CARD_BG, border_color=C_EMERALD)
    tb_c = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p_ct = tf_c.paragraphs[0]
    p_ct.text = "Kesimpulan Utama"
    p_ct.font.size = Pt(20)
    p_ct.font.bold = True
    p_ct.font.color.rgb = C_EMERALD
    p_ct.space_after = Pt(12)

    c_points = [
        "SISNERLING berhasil mentransformasikan pengelolaan neraca lingkungan dari cara manual menjadi sistem digital modern.",
        "Sistem menjamin data yang valid, akuntabel, dan tepat waktu untuk mendukung kebijakan pelestarian lingkungan.",
        "Antarmuka yang ramah pengguna memudahkan adopsi staf tanpa hambatan teknis yang berarti."
    ]
    for cp in c_points:
        p = tf_c.add_paragraph()
        p.text = f"✔  {cp}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_GRAY_LIGHT
        p.space_after = Pt(10)

    add_card(s7, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), C_CARD_BG, border_color=C_AMBER)
    tb_rm = s7.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_rm = tb_rm.text_frame
    tf_rm.word_wrap = True
    p_rmt = tf_rm.paragraphs[0]
    p_rmt.text = "Rencana Pengembangan (Roadmap)"
    p_rmt.font.size = Pt(20)
    p_rmt.font.bold = True
    p_rmt.font.color.rgb = C_AMBER
    p_rmt.space_after = Pt(12)

    rm_points = [
        "Integrasi API Lintas Lembaga: Menghubungkan sistem secara langsung dengan database kementerian/BPS/KLHK.",
        "Modul Neraca Tambahan: Ekspansi cakupan neraca ke sektor mineral, energi, sumber daya air, dan emisi karbon.",
        "Peringatan Dini & Anomali Data: Fitur deteksi cerdas jika terdapat anomali atau lonjakan deplesi sumber daya yang tidak wajar.",
        "Aplikasi Mobile Ringkas: Dashboard versi mobile untuk monitoring ringkas oleh jajaran pimpinan kapan saja."
    ]
    for rp in rm_points:
        p = tf_rm.add_paragraph()
        p.text = f"🚀  {rp}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_GRAY_LIGHT
        p.space_after = Pt(8)

    set_speaker_notes(s7,
        "CATATAN PRESENTER (SPEAKER NOTES):\n"
        "Sebagai penutup, SISNERLING bukan sekadar aplikasi pengolah angka, melainkan fondasi digital untuk tata kelola data lingkungan kita yang lebih baik dan berkelanjutan. "
        "Ke depan, roadmap kami mencakup integrasi API langsung ke instansi eksternal seperti BPS atau KLHK, penambahan modul neraca air dan emisi, serta sistem peringatan dini berbasis anomali data. "
        "Terima kasih atas perhatian bapak/ibu sekalian, saya membuka sesi tanya jawab dan diskusi."
    )

    out_path = r"c:\laragon\www\sisnerling\Sisnerling\_dokumen\Presentasi_SISNERLING_Terbaru.pptx"
    try:
        prs.save(out_path)
        print(f"Presentation updated successfully to {out_path}")
    except Exception as e:
        alt_path = r"c:\laragon\www\sisnerling\Sisnerling\_dokumen\Presentasi_SISNERLING_v2.pptx"
        prs.save(alt_path)
        print(f"Saved to alternative path: {alt_path}")

if __name__ == "__main__":
    create_presentation()
